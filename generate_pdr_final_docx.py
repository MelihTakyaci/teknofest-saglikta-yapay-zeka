#!/usr/bin/env python3
"""
Generate the final TEKNOFEST PDR as a Word document, mapped to the official
table of contents: 1. GIRIS, 2. YONTEM, 3. BULGULAR, 4. SONUC, 5. KAYNAKCA.
Content and all numbers are grounded in the Phase 6 unified results (eşik=0.305).
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path("/Users/melihtakyaci/Documents/TeknofestSagliktaYapayZekâVerisi")
RES = json.load(open(BASE / "reports" / "phase_06_unification" / "phase_06_unified_results.json"))
OUT = BASE / "reports" / "2026_PDR_RaporSon_Final.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def f(x, n=4):
    return f"{x:.{n}f}"


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i].paragraphs[0]
        run = c.add_run(h); run.bold = True; run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]; p.add_run(str(v)).font.size = Pt(9)
    doc.add_paragraph()


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY


def main():
    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(10.5)

    # ── Title block ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("2026 TEKNOFEST SAĞLIKTA YAPAY ZEKÂ YARIŞMASI\nPROJE DETAY RAPORU")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Genetik Varyant Patojenisite Sınıflandırması — Üniversite ve Üzeri Seviyesi\n"
                     "Birleşik Operasyon Noktası: Yığınlanmış Topluluk, eşik = 0.305")
    sr.italic = True; sr.font.size = Pt(10)
    doc.add_paragraph()

    m, p = RES["master"], RES["panels"]
    hu, hp = RES["held_out_union"], RES["held_out_panels"]
    oa, ea = RES["oof_auc"], RES["error_analysis"]

    # ════════════════ 1. GİRİŞ ════════════════
    h1(doc, "1. GİRİŞ")
    doc.add_paragraph(
        "Bu çalışma, ACMG kriterleri ve ClinVar/ClinGen veri tabanları temel alınarak "
        "etiketlenmiş genetik varyantların patojenisitesini, anonimleştirilmiş 353 sütunluk "
        "bir öznitelik şeması üzerinden tahmin eden, klinik güvenliği önceleyen bir makine "
        "öğrenmesi sistemi sunar. Görev, her varyantın Patojenik (1) veya Benign (0) olarak "
        "ikili sınıflandırılmasıdır; değerlendirme F1 Skoru ile yapılır. Sistem; genel (MASTER) "
        "ve üç gen-spesifik panel (Kalıtsal Kanser, Fenilketonüri/PAH, Kistik Fibrozis/CFTR) "
        "üzerinden değerlendirilmektedir.")

    doc.add_heading("1.1 Veri Setleri ve Şema", level=2)
    add_table(doc, ["Veri Seti", "Satır", "Sütun", "Amaç"], [
        ["MASTER (Genel)", "2931", "353", "Genel varyant sınıflandırma"],
        ["KANSER", "388", "353", "Kalıtsal kanser gen paneli"],
        ["PAH", "372", "353", "Fenilketonüri gen paneli"],
        ["CFTR", "111", "353", "Kistik fibrozis gen paneli"],
    ])
    doc.add_paragraph(
        "Öznitelik grupları: AL_1–AL_334 (334 sayısal skor), EK_1–EK_9 (9 birleşik/meta-tahmin "
        "skoru), CAT_1–CAT_6 (6 kategorik), AA_1–AA_2 (2 amino asit), Variant_ID ve Label.")

    doc.add_heading("1.2 Veri Kalitesi: Eksiklik ve Dengesizlik", level=2)
    doc.add_paragraph(
        "MASTER hücrelerinin %54.9'u eksiktir ve eksiklik rastgele değildir (MNAR): patojenik "
        "varyantlarda eksiklik %59.9, benign varyantlarda %41.2'dir; dolayısıyla eksikliğin "
        "kendisi etiket için ayırt edici bilgi taşır. Eğitim setleri patojenik sınıf lehine "
        "dengesizdir (MASTER 2.75:1, KANSER 2.23:1, PAH 5.00:1, CFTR 4.29:1); buna karşın test "
        "setleri dengelidir (1:1). Bu dağılım uyumsuzluğu, eşik kalibrasyonunun dengeli dağılıma "
        "göre yapılmasını zorunlu kılar (bkz. Bölüm 2.3).")

    # ════════════════ 2. YÖNTEM ════════════════
    h1(doc, "2. YÖNTEM")
    doc.add_heading("2.1 Eksiklik-Farkında Özellik Mühendisliği", level=2)
    doc.add_paragraph(
        "MNAR yapısı nedeniyle impütasyon uygulanmamış, eksiklik bir sinyal olarak modellenmiştir: "
        "Variant_ID ve 57 sabit sütun elenmiş; her sayısal öznitelik için ikili eksiklik göstergesi "
        "(miss_*), satır düzeyi toplam eksiklik oranı ve grup düzeyi (AL alt-blokları, EK) eksiklik "
        "sayaçları eklenmiştir. Kategorik öznitelikler tipine göre one-hot / etiket kodlama / ikili "
        "bayrak ile kodlanmıştır. Ağaç tabanlı modeller eksik değerleri doğal olarak işlediğinden "
        "NaN değerleri korunmuş; SMOTE veya yapay örnekleme biyolojik anlamsızlığı nedeniyle "
        "kullanılmamıştır. Nihai matris 655 öznitelik içerir.")

    doc.add_heading("2.2 Model Mimarisi: Küresel Topluluk ve Yığınlama", level=2)
    doc.add_paragraph(
        "MASTER üzerinde eğitilen tek bir küresel topluluk benimsenmiştir: LightGBM (is_unbalance), "
        "XGBoost (scale_pos_weight) ve CatBoost (Balanced sınıf ağırlıkları). Üç modelin örneklem-dışı "
        "(OOF) olasılıkları bir Lojistik Regresyon meta-öğrenicisi ile yığınlanır. Meta-öğrenici "
        "katsayıları (LightGBM 1.47, XGBoost 1.29, CatBoost 2.92) CatBoost'un en yüksek katkıyı "
        "sağladığını gösterir. Tüm aşamalarda sabit tohum (seed=42) ve özdeş 5-katlı StratifiedKFold "
        "bölünmeleri kullanılarak yeniden üretilebilirlik ve sızıntısız panel değerlendirmesi sağlanır.")

    doc.add_heading("2.3 Klinik Maliyetli Eşik Kalibrasyonu ve Operasyon Noktası", level=2)
    doc.add_paragraph(
        "Eğitim dengesiz, test dengeli (1:1) olduğundan eşik yalnızca OOF tahminleri üzerinde ve 1:1 "
        "öncül varsayımıyla kalibre edilmiştir. Patojenik bir varyantı kaçırmanın (yanlış negatif) "
        "klinik maliyeti yanlış alarmdan yüksek olduğundan, maliyet oranı YN:YP = 2:1 alınmış ve "
        "beklenen maliyeti (2·YN + 1·YP) en aza indiren eşik seçilmiştir. İç tutarlılık için, "
        "dağıtılan nihai model olan yığınlanmış topluluğun kalibre eşiği 0.305 tüm panel ve "
        "kohortlarda tek operasyon noktası olarak benimsenmiştir.")

    # ════════════════ 3. BULGULAR ════════════════
    h1(doc, "3. BULGULAR")
    doc.add_heading("3.1 Çapraz Doğrulama ve Topluluk Performansı", level=2)
    add_table(doc, ["Model", "OOF ROC-AUC"], [
        ["LightGBM", f(oa["lightgbm"])], ["XGBoost", f(oa["xgboost"])],
        ["CatBoost", f(oa["catboost"])], ["Yığınlanmış Topluluk", f(oa["stacked"])],
    ])
    doc.add_paragraph(
        f"Yığınlanmış topluluk MASTER üzerinde 0.305 eşiğinde ROC-AUC {f(m['roc_auc'])}, "
        f"PR-AUC {f(m['pr_auc'])}, dengeli F1 {f(m['f1_balanced'])}, duyarlılık {f(m['sensitivity'])}, "
        f"özgüllük {f(m['specificity'])} ve MCC {f(m['mcc'])} elde etmiştir.")

    doc.add_heading("3.2 Panel Bazlı Performans (Sızıntı-Farkında, eşik = 0.305)", level=2)
    add_table(doc, ["Panel", "n", "ROC-AUC", "Dengeli F1", "Duyarlılık", "Özgüllük", "MCC"],
              [[name, p[name]["n"], f(p[name]["roc_auc"]), f(p[name]["f1_balanced"]),
                f(p[name]["sensitivity"]), f(p[name]["specificity"]), f(p[name]["mcc"])]
               for name in ["CFTR", "KANSER", "PAH"]])
    doc.add_paragraph(
        "Paylaşılan varyantlar için örneklem-dışı (OOF) olasılıklar, panele özgü varyantlar için "
        "dondurulmuş topluluk kullanılmıştır. Yığınlama tekil LightGBM'e kıyasla her panelde "
        "ROC-AUC iyileşmesi sağlamıştır (CFTR 0.9402→0.9545, KANSER 0.9141→0.9232, PAH 0.7720→0.7935).")

    doc.add_heading("3.3 Yanlış Pozitif ve Yanlış Negatif Sonuçların Klinik Analizi", level=2)
    doc.add_paragraph(
        f"0.305 operasyon noktasında MASTER örneklem-dışı tahminlerinde {ea['n_fp']} yanlış pozitif "
        f"(YP) ve {ea['n_fn']} yanlış negatif (YN) gözlenmiştir (YN:YP = {ea['fn_fp_ratio']}). YP'nin "
        "YN'yi aşması, modelin 2:1 maliyet politikası gereği bilinçli olarak duyarlılık lehine "
        "yanlılaştırıldığını doğrular. Hataların başlıca biyolojik sürücüsü baskın EK_7 sinyalindeki "
        f"çelişkidir: doğru sınıflandırılan varyantlarda ortalama EK_7 {ea['ek7_mean']['correct']} "
        f"iken kaçırılan patojeniklerde (YN) {ea['ek7_mean']['false_negative']}'e düşmektedir; yani "
        "model, atipik biçimde düşük birleşik hasar skoru taşıyan gerçek patojenikleri kaçırma "
        f"eğilimindedir. Yanlış pozitiflerde EK_7 ortalaması ({ea['ek7_mean']['false_positive']}) "
        "doğru sınıfa yakındır; yükselmiş hasar skoru taşıyan benign varyantlar yanlış alarma yol açar. "
        "İstatistiksel olarak hatalar eksiklik bandlarına göre çift-modlu dağılır: en yüksek mutlak "
        "hata ve YN yoğunluğu %75–100 eksiklik bandındadır (seyrek anotasyon → patojenik kaçırma), "
        "buna karşın iyi anote edilmiş %0–25 bandındaki hatalar sınır-değer varyantların doğasından "
        "kaynaklanır. PAH panelinin zorluğu MASTER örtüşmesinde değil, panele özgü varyantlardaki "
        f"düşük özgüllükte ({f(p['PAH']['specificity'])}) tezahür etmektedir.")
    add_table(doc, ["Eksiklik Bandı", "n", "Hata", "Hata Oranı", "YP", "YN"],
              [[b["band"], b["n"], b["errors"], f(b["error_rate"]), b["fp"], b["fn"]]
               for b in ea["missingness_bands"]])

    doc.add_heading("3.4 EK Ablasyonu ile Veri Sızıntısı Savunması", level=2)
    add_table(doc, ["Yapılandırma", "OOF ROC-AUC"], [
        ["Temel (tüm öznitelikler)", "0.8461"],
        ["Ablasyon (EK_4/5/6/7 yok)", "0.8349"],
        ["Fark (Δ)", "−0.0112"],
    ])
    doc.add_paragraph(
        "0.02'nin altındaki bu düşüş, modelin baskın EK öznitelikilerine düşük bağımlılık "
        "gösterdiğini ve olası etiket sızıntısına (ClinVar döngüselliği) karşı dayanıklı olduğunu "
        "kanıtlar.")

    doc.add_heading("3.5 Ayrılmış Bağımsız Kohort Testi (Dışsal Geçerlilik Simülasyonu)", level=2)
    doc.add_paragraph(
        "Anonimleştirme nedeniyle gerçek anlamda bağımsız bir dış veri seti oluşturulamadığından, "
        "kesin biçimde izole edilmiş bir ayrılmış kohort ile dışsal geçerlilik simülasyonu yapılmıştır. "
        "Temel modeller yalnızca MASTER üzerinde eğitildiğinden, panellerde bulunup MASTER'da "
        "bulunmayan 293 varyant eğitimde hiç görülmemiştir; izolasyon Variant_ID üzerinden kod içinde "
        "doğrulanmıştır (0 sızıntı). Dondurulmuş çıkarım hattı bu kohorta 0.305 eşiğinde uygulanmıştır "
        "(yeniden eğitim yoktur).")
    add_table(doc, ["Kohort", "n", "ROC-AUC", "PR-AUC", "Dengeli F1", "Duyarlılık", "Özgüllük"],
              [["Birleşik", hu["n"], f(hu["roc_auc"]), f(hu["pr_auc"]), f(hu["f1_balanced"]),
                f(hu["sensitivity"]), f(hu["specificity"])]] +
              [[name, hp[name]["n"], f(hp[name]["roc_auc"]), f(hp[name]["pr_auc"]),
                f(hp[name]["f1_balanced"]), f(hp[name]["sensitivity"]), f(hp[name]["specificity"])]
               for name in ["CFTR", "KANSER", "PAH"]])

    doc.add_heading("3.6 Açıklanabilirlik (SHAP)", level=2)
    doc.add_paragraph(
        "SHAP TreeExplainer ile küresel ve yerel açıklamalar üretilmiştir. Ortalama |SHAP| değerine "
        "göre en etkili öznitelikler: EK_7, AL_327, AA_1_R, EK_2, AL_1, AL_16, EK_6, EK_9, AL_311, "
        "AL_319. EK_7'nin baskınlığı, Bölüm 3.3'teki hata analiziyle (düşük EK_7 → kaçırılan patojenik) "
        "tutarlıdır. En yüksek güvenli YP/YN vakaları için yerel öznitelik katkı grafikleri "
        "üretilmiştir (reports/phase_04_optimization/figures/).")

    # ════════════════ 4. SONUÇ ════════════════
    h1(doc, "4. SONUÇ")
    doc.add_paragraph(
        "Önerilen sistem; eksiklik-farkında özellik mühendisliği, klinik maliyete duyarlı örneklem-dışı "
        "eşik kalibrasyonu, çok modelli yığınlama, kanıtlanmış sızıntı izolasyonu ve tek tutarlı "
        f"operasyon noktası (0.305) ile yeniden üretilebilir (seed=42) bir mimaridir. Yığınlanmış "
        f"topluluk OOF ROC-AUC {f(oa['stacked'])} ve ayrılmış kohortta ROC-AUC {f(hu['roc_auc'])} "
        "ile genelleştirilebilir bir performans sergilemektedir.")
    doc.add_heading("4.1 Kısıtlar", level=2)
    for txt in [
        "Dışsal geçerlilik simülasyonu gerçek anlamda bağımsız değildir: ayrılmış kohort aynı anonim "
        "öznitelik çıkarım hattından gelir; izolasyon kesindir ancak bu bir örneklem-dışı tahmindir.",
        f"Ayrılmış kohort ROC-AUC'sinin ({f(hu['roc_auc'])}) iç OOF değerinden ({f(oa['stacked'])}) "
        "yüksek olması 'eğitimden daha iyi' biçiminde yorumlanmamalıdır; neden, görülmemiş CFTR ve "
        "KANSER varyantlarının daha ayrıştırılabilir olmasıdır. Gerçekten zor olan PAH paneli "
        f"ayrılmış kohortta da {f(hp['PAH']['roc_auc'])} ile en düşük değeri verir.",
        "EK döngüselliği tamamen dışlanamaz; anonim öznitelilerin kaynağı doğrulanamadığından "
        "belirsizlik korunur (ablasyon düşük bağımlılık gösterse de).",
        "CFTR örneklem büyüklüğü küçüktür (n=111; ayrılmış alt küme n=34) ve metrikleri yüksek "
        "varyans taşır.",
        f"PAH paneli düşük özgüllük ({f(p['PAH']['specificity'])}) ile en zayıf performansı verir; "
        "5:1 dengesizlik ve yüksek eksiklik temel nedendir ve öncelikli iyileştirme hedefidir.",
    ]:
        doc.add_paragraph(txt, style="List Bullet")
    doc.add_paragraph(
        "Gelecek çalışmalar: Optuna ile hiperparametre optimizasyonu, panel-bazlı (leave-one-panel-out) "
        "eşik kalibrasyonu ve daha derin hata analizi mevcut altyapı üzerinde doğrudan uygulanabilir.")

    # ════════════════ 5. KAYNAKÇA ════════════════
    h1(doc, "5. KAYNAKÇA")
    refs = [
        "Richards S, et al. Standards and guidelines for the interpretation of sequence variants: "
        "a joint consensus recommendation of the ACMG and AMP. Genetics in Medicine. 2015;17(5):405-424.",
        "Landrum MJ, et al. ClinVar: improving access to variant interpretations and supporting "
        "evidence. Nucleic Acids Research. 2018;46(D1):D1062-D1067.",
        "Ke G, et al. LightGBM: A Highly Efficient Gradient Boosting Decision Tree. Advances in "
        "Neural Information Processing Systems (NeurIPS). 2017.",
        "Chen T, Guestrin C. XGBoost: A Scalable Tree Boosting System. Proceedings of the 22nd ACM "
        "SIGKDD International Conference on Knowledge Discovery and Data Mining. 2016:785-794.",
        "Prokhorenkova L, et al. CatBoost: unbiased boosting with categorical features. Advances in "
        "Neural Information Processing Systems (NeurIPS). 2018.",
        "Lundberg SM, Lee SI. A Unified Approach to Interpreting Model Predictions. Advances in "
        "Neural Information Processing Systems (NeurIPS). 2017.",
        "Karczewski KJ, et al. The mutational constraint spectrum quantified from variation in "
        "141,456 humans (gnomAD). Nature. 2020;581:434-443.",
    ]
    for i, ref in enumerate(refs, 1):
        doc.add_paragraph(f"[{i}] {ref}")

    doc.save(OUT)
    print(f"Saved -> {OUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
